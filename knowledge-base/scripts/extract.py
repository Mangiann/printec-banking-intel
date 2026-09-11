#!/usr/bin/env python3
"""Deterministic document extractor for the Printec knowledge-base.

Converts a dropped document (pptx / xlsx / docx / pdf / txt-like) into plain
markdown text on stdout, so the nightly data-dump processor (an LLM agent) can
digest the *text* of a 60 MB deck without ever loading the binary into context.

Usage:
    python extract.py "<path-to-file>"                 # -> markdown text on stdout
    python extract.py "<path-to-file>" --meta-only      # -> JSON metadata only
    python extract.py "<path-to-file>" --out "<file>"   # -> write text to a file

Design notes:
  * Pure extraction, no judgement — the LLM does tagging/summary/routing.
  * Office temp/lock files ("~$*") are refused (they are not real documents).
  * stdout is forced to UTF-8 so Windows consoles don't crash on non-ASCII
    (the same gotcha that once broke verify.py).
  * Never raises for content problems: emits a "[extract error: ...]" marker so
    the caller still gets a usable (if partial) result and the run continues.
"""
import argparse
import hashlib
import json
import os
import sys

# Force UTF-8 stdout/stderr (Windows cp1252 consoles otherwise crash on e.g. €, →).
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

MAX_XLSX_ROWS = 400      # cap per sheet so a giant export can't blow up the digest
MAX_XLSX_COLS = 40


def sha256_of(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def is_lock_file(path):
    return os.path.basename(path).startswith("~$")


# ---------------------------------------------------------------- pptx
def extract_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n## Slide {i}")
        for shape in slide.shapes:
            try:
                if shape.has_table:
                    tbl = shape.table
                    rows = list(tbl.rows)
                    if not rows:
                        continue
                    def cells(r):
                        return [c.text.strip().replace("\n", " ") for c in r.cells]
                    header = cells(rows[0])
                    out.append("| " + " | ".join(header) + " |")
                    out.append("| " + " | ".join("---" for _ in header) + " |")
                    for r in rows[1:]:
                        out.append("| " + " | ".join(cells(r)) + " |")
                    continue
            except Exception:
                pass
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                txt = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if txt:
                    out.append(txt)
        # speaker notes
        try:
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    out.append(f"_Notes:_ {notes}")
        except Exception:
            pass
    return "\n".join(out).strip()


# ---------------------------------------------------------------- xlsx
def extract_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n## Sheet: {ws.title}")
        n = 0
        header_done = False
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).replace("\n", " ").strip()
                    for v in row[:MAX_XLSX_COLS]]
            if not any(vals):
                continue
            out.append("| " + " | ".join(vals) + " |")
            if not header_done:
                out.append("| " + " | ".join("---" for _ in vals) + " |")
                header_done = True
            n += 1
            if n >= MAX_XLSX_ROWS:
                out.append(f"_… truncated at {MAX_XLSX_ROWS} rows_")
                break
    wb.close()
    return "\n".join(out).strip()


# ---------------------------------------------------------------- docx
def extract_docx(path):
    import docx
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in r.cells) + " |")
    return "\n".join(out).strip()


# ---------------------------------------------------------------- pdf
def extract_pdf(path):
    try:
        from pypdf import PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader  # older fallback
        except Exception:
            return ("[extract error: no PDF library installed — run "
                    "`pip install pypdf`, or read this PDF with the Read tool]")
    reader = PdfReader(path)
    out = []
    for i, page in enumerate(reader.pages, 1):
        try:
            txt = (page.extract_text() or "").strip()
        except Exception as e:
            txt = f"[page {i} extract error: {e}]"
        if txt:
            out.append(f"\n## Page {i}\n{txt}")
    return "\n".join(out).strip()


TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".log"}


def extract_text(path):
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read().strip()


EXTRACTORS = {
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
}


def extract(path):
    ext = os.path.splitext(path)[1].lower()
    fn = EXTRACTORS.get(ext)
    if fn is None and ext in TEXT_EXTS:
        fn = extract_text
    if fn is None:
        return f"[extract error: unsupported file type '{ext}']"
    try:
        return fn(path) or "[extract warning: no extractable text found]"
    except Exception as e:
        return f"[extract error: {type(e).__name__}: {e}]"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--meta-only", action="store_true")
    ap.add_argument("--out")
    args = ap.parse_args()

    path = args.path
    if not os.path.isfile(path):
        print(json.dumps({"ok": False, "error": "not a file", "file": path}))
        sys.exit(2)
    if is_lock_file(path):
        print(json.dumps({"ok": False, "error": "office lock/temp file (~$)", "file": path}))
        sys.exit(3)

    meta = {
        "file": os.path.basename(path),
        "ext": os.path.splitext(path)[1].lower(),
        "bytes": os.path.getsize(path),
        "sha256": sha256_of(path),
        "ok": True,
    }
    if args.meta_only:
        print(json.dumps(meta, ensure_ascii=False))
        return

    text = extract(path)
    header = (f"<!-- extracted: {meta['file']} | {meta['ext']} | "
              f"{meta['bytes']} bytes | sha256:{meta['sha256'][:12]} -->\n")
    body = header + text
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(body)
        print(json.dumps({**meta, "out": args.out, "chars": len(text)}, ensure_ascii=False))
    else:
        sys.stdout.write(body)


if __name__ == "__main__":
    main()
