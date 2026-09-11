"""Deterministic consolidation - group a run's candidates by ENTITY (registrable domain) into a primary
source + its child endpoints, so a reviewer navigates *entities*, not *pages*.

Why this is a separate step: the scout runs many cells in PARALLEL with no shared state, so one entity
(e.g. Bank of Greece) is emitted as many rows by many cells. No single agent can merge them - only a
post-step with the full list can. This module groups by registrable domain, elects a primary "hub"
endpoint per group (shallowest path, then best score), and lists the rest as child endpoints - WITHOUT
dropping any candidate (grouping is a VIEW; nothing is lost).

Scope: structural grouping only. JUDGMENT (keep / evidence / pointer / marginal) is NOT done here - that
belongs to the scout at capture and the pre-labeler at review. Pure, total, no I/O.
"""
import re
from urllib.parse import urlparse

from _safe import dget

# multi-label public suffixes we must not split (so eprocurement.gov.gr -> that, not gov.gr)
_MULTI_SUFFIX = {"gov.gr", "com.gr", "net.gr", "org.gr", "edu.gr", "co.gr",
                 "co.uk", "org.uk", "gov.uk", "com.cy", "gov.cy", "org.cy"}
# shared institutional domains where each SUBDOMAIN is a distinct entity (TED, Cohesion, EPRS all live on
# europa.eu but are different bodies) - group these by full hostname, not registrable domain.
_GROUP_BY_HOST = {"europa.eu"}
_NOKEY = "__nokey__:"


def _hostname(url):
    try:
        host = (urlparse(url).hostname or "").lower()
    except (ValueError, TypeError):
        return ""
    return host[4:] if host.startswith("www.") else host


def registrable_domain(url):
    """eTLD+1 (best-effort, no public-suffix list) - the entity key. Fail-closed to '' on a bad URL."""
    host = _hostname(url)
    if not host:
        return ""
    labels = [x for x in host.split(".") if x]
    if len(labels) <= 2:
        return ".".join(labels)
    return ".".join(labels[-3:]) if ".".join(labels[-2:]) in _MULTI_SUFFIX else ".".join(labels[-2:])


def entity_key(url):
    """The grouping key: the registrable domain, EXCEPT for shared institutional domains whose subdomains
    are distinct entities (europa.eu) - there the full hostname is the entity."""
    reg = registrable_domain(url)
    return _hostname(url) if reg in _GROUP_BY_HOST else reg


def _path_depth(url):
    try:
        return len([s for s in (urlparse(url).path or "").split("/") if s])
    except (ValueError, TypeError):
        return 99


def _score(c):
    def n(v):
        return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
    return n(c.get("credibility_score")) + n(c.get("usefulness_score"))


def _entity_name(primary):
    name = primary.get("source_name")
    if isinstance(name, str) and name.strip():
        base = re.split(r"\s[—–-]\s|[:(]", name.strip())[0].strip()
        return base or name.strip()
    return registrable_domain(primary.get("canonical_url") or "") or "unknown"


def group_candidates(candidates):
    """Group by registrable domain -> list of groups (biggest first). Each group:
    {entity, entity_name, primary_id, member_ids (primary first), count}. Total: non-dict / no-candidate_id
    entries are skipped; a candidate with no usable domain becomes its own singleton group."""
    groups = {}
    for c in candidates:
        if not isinstance(c, dict):
            continue
        cid = c.get("candidate_id")
        if not isinstance(cid, str) or not cid:
            continue
        key = entity_key(c.get("canonical_url") or "") or (_NOKEY + cid)
        groups.setdefault(key, []).append(c)

    out = []
    for key, members in groups.items():
        # primary = shallowest path (the hub), then highest (credibility+usefulness), then candidate_id
        ordered = sorted(members, key=lambda c: (_path_depth(c.get("canonical_url") or ""),
                                                 -_score(c), str(c.get("candidate_id"))))
        primary = ordered[0]
        out.append({
            "entity": "" if key.startswith(_NOKEY) else key,
            "entity_name": _entity_name(primary),
            "primary_id": primary["candidate_id"],
            "member_ids": [m["candidate_id"] for m in ordered],
            "count": len(ordered),
        })
    out.sort(key=lambda g: (-g["count"], g["entity_name"].lower(), g["primary_id"]))
    return out


def write_grouped_workbook(candidates, groups, out_path, *, role_by_id=None):
    """Render the grouping to an .xlsx: one section per entity (a bold PRIMARY row + its endpoint rows).
    A VIEW for navigation - nothing is decided or dropped here. If `role_by_id` (candidate_id -> source-role
    like durable_source/evidence/pointer/marginal) is given, a `kind` column is filled. Returns out_path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    roles = role_by_id if isinstance(role_by_id, dict) else {}
    by_id = {c["candidate_id"]: c for c in candidates
             if isinstance(c, dict) and isinstance(c.get("candidate_id"), str)}
    HEADERS = ["group", "role", "kind", "entity", "priority", "source_class", "source_subtype",
               "source_name", "canonical_url", "url_verified", "why_relevant"]
    wb = Workbook(); ws = wb.active; ws.title = "Grouped"
    for i, h in enumerate(HEADERS, 1):
        ws.cell(1, i, h).font = Font(bold=True)
    pfill, cfill = PatternFill("solid", fgColor="DCE6F1"), PatternFill("solid", fgColor="F2F2F2")

    def _cell(v):
        return v if v is None or isinstance(v, (str, int, float, bool)) else str(v)

    r = 2
    for gno, g in enumerate(groups, 1):
        for j, cid in enumerate(g["member_ids"]):
            c = by_id.get(cid) or {}
            primary = j == 0
            name = c.get("source_name") if isinstance(c.get("source_name"), str) else ""
            vals = [gno, "PRIMARY" if primary else "endpoint", roles.get(cid, ""),
                    g["entity_name"] if primary else "",
                    c.get("priority"), c.get("source_class"), c.get("source_subtype"),
                    name if primary else "   └ " + name, c.get("canonical_url"),
                    dget(c.get("url_verification"), "status"), dget(c.get("discovery_basis"), "why_relevant")]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(r, i, _cell(v))
                cell.fill = pfill if primary else cfill
                if primary:
                    cell.font = Font(bold=True)
            r += 1
    widths = {"entity": 30, "source_name": 48, "canonical_url": 50, "why_relevant": 50, "kind": 15}
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path


def summarize(groups):
    total_members = sum(g["count"] for g in groups)
    multi = [g for g in groups if g["count"] > 1]
    return {
        "candidates": total_members,
        "entities": len(groups),
        "multi_endpoint_entities": len(multi),
        "rows_folded": total_members - len(groups),   # how many rows collapse under a primary
    }
