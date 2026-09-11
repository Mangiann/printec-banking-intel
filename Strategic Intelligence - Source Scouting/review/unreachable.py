"""The "check by hand" list: sources the automatic recovery ladder could not reach, for a human to resolve.

After a round has run every rung it can (HEAD, GET, WAF handshake, web search, real Chrome), a few sources
are still unreachable — behind a login, a CAPTCHA, geo-blocked, or genuinely gone. Rather than guess, the
round hands them to the human as a plainly-named Excel per country, with the machine facts locked and a small
set of yellow boxes to fill: the working URL (if they find one), a status, and a note. `apply_manual_url_checks`
reads it back and folds their answers in exactly like any other recovery rung.

Deterministic, total on malformed on-disk records (candidates load with no schema enforcement — same posture
as review/excel.py).
"""
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from _safe import dget, sstr

_SHEET = "Check by hand"
# machine columns (locked, information only) + human columns (yellow, the only editable ones)
MACHINE = ["source_name", "workstream", "why_relevant", "url_we_tried", "last_status", "what_we_tried",
           "run_id", "candidate_id"]
HUMAN = ["working_url", "status", "notes"]
HEADERS = MACHINE + HUMAN
STATUS_VOCAB = ["works", "moved", "dead", "needs_login", "not_relevant"]
_EDITABLE = set(HUMAN)


def _run_prefix(registry_db):
    """GR.db -> 'dr_gr_' — a country's runs share this candidate/run-id prefix (mirrors master_export)."""
    stem = Path(registry_db).stem
    return f"dr_{stem.lower()}_"


def gather_unreachable(data_root, *, run_prefix="", role_dir=None):
    """Durable sources whose URL is STILL 'failed' after all recovery, deduped by canonical_url.

    A candidate is included iff (a) its run's role verdict is 'keep' (a durable monitoring source, not a
    one-off article or a directory) and (b) its url_verification is 'failed'. run_prefix scopes to a
    country's runs; "" means all runs."""
    data_root = Path(data_root)
    by_url = {}
    for cf in sorted((data_root / "source_candidates").glob("*.jsonl")):
        run = cf.stem
        if run_prefix and not run.startswith(run_prefix):
            continue
        ai = data_root / "run_logs" / run / "agent_inputs"
        roles = _load_map(ai / "role_verdicts.json", "candidate_id", "verdict")
        statuses = _load_map(ai / "remaining_failures.json", "candidate_id", "http_status")
        for line in cf.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            try:
                c = json.loads(line)
            except ValueError:
                continue
            if not isinstance(c, dict):
                continue
            cid = sstr(c.get("candidate_id"))
            if roles.get(cid) != "keep":
                continue
            if dget(c.get("url_verification"), "status") != "failed":
                continue
            url = sstr(c.get("canonical_url")) or sstr(c.get("source_url"))
            key = url or cid
            if key in by_url:
                continue
            wss = c.get("workstreams") or ([c.get("workstream")] if c.get("workstream") else [])
            by_url[key] = {
                "source_name": sstr(c.get("source_name")),
                "workstream": ",".join(x for x in wss if isinstance(x, str)),
                "why_relevant": sstr(dget(c.get("discovery_basis"), "why_relevant")),
                "url_we_tried": url,
                "last_status": sstr(statuses.get(cid)) or "unreachable",
                "what_we_tried": "HEAD, GET, WAF-handshake, web search, real Chrome — all failed",
                "run_id": run,
                "candidate_id": cid,
            }
    return sorted(by_url.values(), key=lambda r: (r["run_id"] or "", r["source_name"] or ""))


def _load_map(path, key, val):
    out = {}
    if not Path(path).exists():
        return out
    try:
        rows = json.loads(Path(path).read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return out
    for r in rows if isinstance(rows, list) else []:
        if isinstance(r, dict) and sstr(r.get(key)) is not None:
            out[sstr(r.get(key))] = r.get(val)
    return out


def export_unreachable_workbook(data_root, out_path, *, registry_db=None, run_prefix=None, country_name=""):
    """Write the check-by-hand workbook. Returns the row count. Writes even when 0 (a clean 'nothing to do')."""
    prefix = run_prefix if run_prefix is not None else (_run_prefix(registry_db) if registry_db else "")
    rows = gather_unreachable(data_root, run_prefix=prefix)
    _build(out_path, rows, country_name)
    return len(rows)


def _build(out_path, rows, country_name):
    wb = Workbook()
    info = wb.active
    info.title = "READ ME FIRST"
    intro = [
        ("SOURCES TO CHECK BY HAND", ""),
        ("What this is", f"Sources for {country_name or 'this country'} our automatic checks could NOT reach "
                         "(login walls, CAPTCHAs, geo-blocks, or genuinely gone). We tried everything: a plain "
                         "fetch, a browser-style fetch, web search, and a real Chrome browser."),
        ("", ""),
        ("Step 1", "Open the 'Check by hand' tab. Each row is one source, with the link we tried in "
                   "'url_we_tried'."),
        ("Step 2", "Open that link yourself in your browser."),
        ("Step 3", "Fill the yellow boxes: put the address that works in 'working_url' (if the page moved, "
                   "paste the NEW address), and choose a 'status'."),
        ("status = works", "the link is fine as-is (you can leave working_url empty)."),
        ("status = moved", "it lives at a new address — paste it in working_url."),
        ("status = dead", "the source is genuinely gone."),
        ("status = needs_login", "it's real but needs a login/subscription — we'll set it aside for manual "
                                 "upload later."),
        ("status = not_relevant", "not actually a banking/payments source we want."),
        ("Step 4", "Save the file and hand it back. Only the yellow columns are read; the grey ones are "
                   "information only."),
        ("", ""),
        ("Sources to check", len(rows)),
    ]
    for r, (k, v) in enumerate(intro, 1):
        info.cell(row=r, column=1, value=k).font = Font(bold=True)
        info.cell(row=r, column=2, value=v)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 104

    ws = wb.create_sheet(_SHEET)
    hdr_font, hdr_fill = Font(bold=True), PatternFill("solid", fgColor="DDDDDD")
    edit_fill = PatternFill("solid", fgColor="FFF7CC")
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = hdr_font
        cell.fill = edit_fill if h in _EDITABLE else hdr_fill
    for r, row in enumerate(rows, start=2):
        for i, h in enumerate(HEADERS, 1):
            cell = ws.cell(row=r, column=i, value=row.get(h))
            if h in _EDITABLE:
                cell.protection = Protection(locked=False)
                cell.fill = edit_fill
    if rows:
        col = get_column_letter(HEADERS.index("status") + 1)
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(STATUS_VOCAB), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
    widths = {"source_name": 34, "why_relevant": 52, "url_we_tried": 50, "working_url": 50, "what_we_tried": 44,
              "notes": 30, "candidate_id": 26, "run_id": 20}
    for i, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    ws.freeze_panes = "A2"
    ws.protection.sheet = True
    wb.save(out_path)


def import_unreachable_checks(xlsx_path):
    """Read a returned workbook -> {"checks": [...], "errors": [...]}. TOTAL: a bad file -> an errors result,
    never a raise. Reads ONLY rows the human actually filled (a status set)."""
    try:
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return {"checks": [], "errors": [f"could not open workbook: {e}"]}
    if _SHEET not in wb.sheetnames:
        wb.close()
        return {"checks": [], "errors": [f"no '{_SHEET}' sheet"]}
    ws = wb[_SHEET]
    data = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    if not data:
        return {"checks": [], "errors": ["empty sheet"]}
    head = [sstr(x) and str(x).strip() for x in data[0]]
    idx = {h: i for i, h in enumerate(head) if h}
    for req in ("candidate_id", "status"):
        if req not in idx:
            return {"checks": [], "errors": [f"missing column '{req}'"]}

    def cell(row, name):
        i = idx.get(name)
        return row[i] if i is not None and i < len(row) else None

    checks, errors = [], []
    for row in data[1:]:
        status = sstr(cell(row, "status"))
        if not status:
            continue                                       # unfilled row -> skip
        status = status.strip().lower()
        cid = sstr(cell(row, "candidate_id"))
        if not cid:
            errors.append("a filled row has no candidate_id")
            continue
        if status not in STATUS_VOCAB:
            errors.append(f"{cid}: status '{status}' not one of {STATUS_VOCAB}")
            continue
        working = sstr(cell(row, "working_url"))
        if status == "moved" and not working:
            errors.append(f"{cid}: status 'moved' needs a working_url")
            continue
        checks.append({"candidate_id": cid, "run_id": sstr(cell(row, "run_id")), "status": status,
                       "working_url": working, "notes": sstr(cell(row, "notes"))})
    return {"checks": checks, "errors": errors}
