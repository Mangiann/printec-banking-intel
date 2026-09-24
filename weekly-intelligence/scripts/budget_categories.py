#!/usr/bin/env python3
"""
budget_categories.py — the 2027 budget in the categories of the Printec budget Excel.

The Excel (budget/budget_clean.xlsx) budgets each country by category (HW, SW, SV, OUT) and
subcategory (ATM, EFT-POS, AML ...). The 2027 numbers (budget/budget_2027_board.json) are set by
country and product line. This step splits every cell's 2027 numbers (lowest acceptable, budget,
maximum, and the reviewed number before the plan) into the four categories and their subcategories,
so the categories add up exactly to the plan and every euro can be traced to what produces it.

Allocation rules (decided with the user on 17/09/2026; nothing here re-estimates a number):
  * The cell's own 2026 split across HW / SW / SV / OUT (from the Excel) is the deal mix. A dated
    order or a share win is split the way that country's revenue for the subcategory already splits,
    so installation and first-year support land in SV through the Excel's own structure.
  * The model's recurring growth goes to SV and OUT in their 2026 ratio; the model's named-deal and
    market-outlook growth follows the full 2026 mix.
  * Hires go to SV (to OUT when the line is managed services); capability decisions go to SW; price
    and efficiency steps go to SV and OUT in their 2026 ratio.
  * The three closing facts are orders: full mix. Any balance left after the layers (e.g. the
    maximum case's extra new business) follows the full mix.
  * A cell with no 2026 revenue (new business) uses the group's 2026 mix for its line, except managed
    services which is 100% OUT. Its subcategory is the line's lead subcategory (managed services: ATM).
  * Retail lines and "Other" rows (no product name) are carried flat inside their own category; no separate bucket.
  * Profit holds the 2026 margin of each country x category x subcategory (fallbacks: the country's
    category margin, the group's subcategory margin, the group's category margin).

Writes budget/budget_2027_categories.json (embedded in the dashboard's encrypted budget payload) and
budget/budget_2027_by_category.xlsx (the 2026 Excel's layout, with the 2027 columns).
"""
import json, os, sys, datetime
from collections import defaultdict, OrderedDict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
CATS = ["HW", "SW", "SV", "OUT"]
CAT_LABEL = {"HW": "Hardware", "SW": "Software", "SV": "Services", "OUT": "Outsourcing"}
SCEN = OrderedDict([("floor", "floor"), ("budget", "target"), ("upside", "upside")])   # json key -> board key
# where a line lands when a country has no 2026 row for it
LEAD_SUB = {"atm_recycling": "ATM", "pos_acquiring": "EFT-POS", "self_service": "Self-Service Kiosks",
            "compliance_resilience": "Other Security & Compliance", "fraud_monitoring": "Fraud Prevention & Detection",
            "digital_onboarding": "eSignature", "managed_services": "ATM", "card_issuing": "Cards & Loyalty",
            "physical_security": "ATM & Safe Security solutions", "core_digital": "eCommerce",
            "payments_instant": "Other e-payments", "payment_security_hsm": "Other Security & Compliance"}
KIND_LABEL = {"date": "dated order or tender", "share": "share win", "hire": "hire", "capability": "capability decision",
              "pricing": "price step", "efficiency": "efficiency step"}

def load(p):
    with open(p, encoding="utf-8") as f:
        return json.load(f)

def save(p, obj):
    with open(p, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)

def read_budget(path, bmap):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    recs, country, cat, order = [], None, None, []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        n = str(r[0]).strip()
        if n.startswith("Printec"):
            country, cat = n, None
            if n not in order: order.append(n)
            continue
        if n in CATS:
            cat = n; continue
        if n == "TOTAL" or n.startswith("Source") or country is None or cat is None:
            continue
        recs.append({"entity": country, "code": bmap["countries"].get(country), "cat": cat, "sub": n,
                     "rev": float(r[1] or 0), "rp": float(r[2] or 0)})
    return recs, order

def norm(d):
    t = sum(v for v in d.values() if v > 0)
    return {k: (v / t if t > 0 and v > 0 else 0.0) for k, v in d.items()} if t > 0 else {}

def M(v):
    return "EUR %.2fm" % (v / 1e6)

def compute(root):
    bdir = os.path.join(root, "budget")
    bmap = load(os.path.join(root, "weekly-intelligence", "references", "budget_map.json"))
    labels = {c["id"]: c["label"] for c in load(os.path.join(root, "weekly-intelligence", "references", "product_map.json"))["categories"]}
    B = load(os.path.join(bdir, "budget_2027_board.json"))
    single = bool(B.get("single_number"))
    scen = OrderedDict([("budget", "target")]) if single else SCEN
    ST = load(os.path.join(bdir, "budget_2027_stretch.json")) if os.path.exists(os.path.join(bdir, "budget_2027_stretch.json")) else None
    st_val = {}
    if ST:
        scen = OrderedDict(list(scen.items()) + [("stretch", "stretch")])
        for x in ST.get("cells", []): st_val["%s|%s" % (x["code"], x["line"])] = x["stretch"]
    recs, entity_order = read_budget(os.path.join(bdir, "budget_clean.xlsx"), bmap)
    retail = set(bmap["retail"])
    name = {}
    for c in B["cells"] + B.get("new_business", []):
        name[c["code"]] = c["country"]
    entity_of = {v: k for k, v in bmap["countries"].items()}

    # ---- 2026 by cell x category x subcategory; retail/other by country x category x subcategory ----
    base = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))      # (cc,line) -> cat -> sub -> rev
    base_rp = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    flat = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: [0.0, 0.0])))  # cc -> cat -> sub -> [rev, rp]
    flat_kind = {}
    grp_line = defaultdict(lambda: defaultdict(float))                       # line -> cat -> rev (group)
    grp_sub_rm = defaultdict(lambda: [0.0, 0.0])                             # (cat,sub) -> [rev,rp]
    grp_cat_rm = defaultdict(lambda: [0.0, 0.0])
    for r in recs:
        cc = r["code"]
        if not cc:
            continue
        line = bmap["lines"].get(r["sub"])
        if r["cat"] == "OUT" and line is None and r["sub"] not in retail:
            line = "managed_services"
        grp_sub_rm[(r["cat"], r["sub"])][0] += r["rev"]; grp_sub_rm[(r["cat"], r["sub"])][1] += r["rp"]
        grp_cat_rm[r["cat"]][0] += r["rev"]; grp_cat_rm[r["cat"]][1] += r["rp"]
        if line is None:
            # retail lines and budget with no product name: carried flat inside their own category, no separate
            # bucket (decisions of 17/09/2026: first Other, then retail)
            line = "retail_flat" if r["sub"] in retail else "other_flat"
        base[(cc, line)][r["cat"]][r["sub"]] += r["rev"]
        base_rp[(cc, line)][r["cat"]][r["sub"]] += r["rp"]
        grp_line[line][r["cat"]] += r["rev"]

    def rm_for(cc, line, cat, sub):
        """2026 margin to hold: cell's own sub, then the country's category, then the group's sub, then the group's category."""
        rv, rp = base[(cc, line)][cat].get(sub, 0.0), base_rp[(cc, line)][cat].get(sub, 0.0)
        if rv > 0: return rp / rv
        crv = sum(sum(s.values()) for (c2, l2), cats in base.items() if c2 == cc for k, s in cats.items() if k == cat)
        crp = sum(sum(s.values()) for (c2, l2), cats in base_rp.items() if c2 == cc for k, s in cats.items() if k == cat)
        if crv > 0: return crp / crv
        g = grp_sub_rm.get((cat, sub))
        if g and g[0] > 0: return g[1] / g[0]
        g = grp_cat_rm.get(cat)
        return g[1] / g[0] if g and g[0] > 0 else 0.5

    # ---- allocate every cell ----
    # Small combinations (each under EUR 50k) that the plan file does not carry as cells are held at the
    # model's estimate in every scenario, which is how the plan's totals count them.
    cells_out = []
    all_cells = B["cells"] + B.get("new_business", [])
    for (cc, ln), cats in list(base.items()):
        if ln in ("other_flat", "retail_flat"):
            tot = sum(v for d in cats.values() for v in d.values())
            all_cells.append({"code": cc, "country": name.get(cc, entity_of.get(cc, cc)), "line": ln, "target_2026": tot, "new_business": False, "flat_other": True,
                              "board": {"target": tot, "floor": tot, "upside": tot, "today": {"target": tot, "floor": tot, "upside": tot}, "parts": {}, "plan_actions": []}})
    labels["other_flat"] = "no product name, carried flat"; labels["retail_flat"] = "retail, carried flat"

    have = {c["code"] + "|" + c["line"] for c in all_cells}
    if ST:
        for x in ST.get("new_cells", []):
            k = x["code"] + "|" + x["line"]
            if k in have: continue
            have.add(k); st_val[k] = x["stretch"]
            all_cells.append({"code": x["code"], "country": x["country"], "line": x["line"], "target_2026": 0.0, "new_business": True, "stretch_only": True,
                              "board": {"target": 0.0, "floor": 0.0, "upside": 0.0, "today": {"target": 0.0}, "parts": {"case": "new"}, "plan_actions": []}})
    A26 = load(os.path.join(bdir, "budget_actions.json"))
    P27 = load(os.path.join(bdir, "budget_2027.json"))
    pk = {c["code"] + "|" + c["line"]: c for c in P27.get("cells", [])}
    for c in A26["cells"]:
        k = c["code"] + "|" + c["line"]
        if k in have: continue
        est = (pk.get(k) or {}).get("base", {}).get("target", c["target"])
        all_cells.append({"code": c["code"], "country": c["country"], "line": c["line"], "target_2026": c["target"], "new_business": False, "small": True,
                          "board": {"target": est, "floor": est, "upside": est, "today": {"target": est, "floor": est, "upside": est}, "parts": {}, "plan_actions": []}})
    for c in all_cells:
        cc, line = c["code"], c["line"]
        bd = c["board"]
        b26 = {k: sum(base[(cc, line)][k].values()) for k in CATS}
        tot26 = sum(b26.values())
        if tot26 > 0:
            mix_full = norm(b26)
        elif line == "managed_services":
            mix_full = {"OUT": 1.0}
        else:
            mix_full = norm({k: grp_line[line].get(k, 0.0) for k in CATS}) or {"SW": 0.6, "SV": 0.4}
        mix_rec = norm({k: b26[k] for k in ("SV", "OUT")}) or ({"OUT": 1.0} if line == "managed_services" else {"SV": 1.0})
        parts = bd.get("parts") or {}
        psum = sum(abs(parts.get(k, 0.0) or 0.0) for k in ("recurring", "named_deals_net", "market_outlook"))
        rec_share = (abs(parts.get("recurring", 0.0) or 0.0) / psum) if psum > 0 else 0.0
        hire_mix = {"OUT": 1.0} if line == "managed_services" else {"SV": 1.0}
        today = bd.get("today") or {}
        out = {"code": cc, "country": c["country"], "line": line, "line_label": labels.get(line, line),
               "new_business": bool(c.get("new_business")), "y2026": {k: round(b26[k], 2) for k in CATS},
               "mix": {k: round(v, 4) for k, v in mix_full.items()}, "scen": {}}
        for sk, bk in list(scen.items()) + [("today", "target")]:
            alloc = defaultdict(float); src = defaultdict(list)
            def put(mix, eur, cat_label, label):
                if not eur: return
                for k, w in mix.items():
                    if w <= 0: continue
                    alloc[k] += eur * w; src[k].append({"label": label, "eur": round(eur * w, 2), "kind": cat_label})
            if sk == "today":
                T = today.get("target", bd.get("target")) if today else bd.get("target")
            elif sk == "stretch":
                T = st_val.get(cc + "|" + line, bd.get("target"))
            else:
                T = bd.get(bk)
            # a new-business combination that already has 2026 revenue in the Excel: the plan's number is new
            # revenue on top of it, so the 2026 revenue is carried flat underneath (this is how the plan's totals count it)
            if c.get("new_business") and tot26 > 0:
                T = (T or 0.0) + tot26
            t_s = (today.get(bk) if today else None)
            if t_s is None: t_s = c.get("target_2026", 0.0)
            # layer A: the model's growth to the reviewed number before the plan
            A = (t_s or 0.0) - (c.get("target_2026") or 0.0)
            if A:
                put(mix_rec, A * rec_share, "model_recurring", "recurring revenue growth in the model")
                put(mix_full, A * (1 - rec_share), "model_deals", "named deals and market outlook in the model")
            if sk != "today":
                # layer B: the plan's actions
                for a in bd.get("plan_actions", []):
                    eur = a.get("eur_floor" if sk == "floor" else "eur_budget") or 0.0
                    if not eur: continue
                    kind = a.get("kind", "date")
                    title = a.get("deal_title") or ""
                    who = title.split(" — ")[0].split(" - ")[0].strip()
                    if not who or who == c["country"] or len(who) < 3:
                        who = (title or a.get("what", "")).strip()
                    who = who[:80].rstrip(" .")
                    if kind == "date": lab = "the dated order or tender at " + who if title else "the dated order (" + who + ")"
                    elif kind == "share": lab = "the share win at " + who if title else "the share win (" + who + ")"
                    elif kind == "hire": lab = "the hire (" + who[0].lower() + who[1:] + ")"
                    elif kind == "capability": lab = "the capability decision (" + who[0].lower() + who[1:] + ")"
                    else: lab = "the " + KIND_LABEL.get(kind, kind) + " (" + who[0].lower() + who[1:] + ")"
                    if kind in ("date", "share"): put(mix_full, eur, kind, lab)
                    elif kind == "hire": put(hire_mix, eur, kind, lab)
                    elif kind == "capability": put({"SW": 1.0}, eur, kind, lab)
                    else: put(mix_rec, eur, kind, lab)
                for a in bd.get("staffing", []):
                    put(hire_mix, a.get("eur_budget") or 0.0, "capacity", "delivery by existing staff (" + (a.get("what", "")[:70]).rstrip(" .") + ")")
                # layer C: the closing facts (floor only)
                if sk == "floor" and bd.get("gap_closer_eur"):
                    put(mix_full, bd["gap_closer_eur"], "closing_fact", "one of the three closing facts in the floor")
            # layer D: the balance; for the stretch, what sits above the budget is the stretch layer
            got = tot26 + sum(alloc.values())
            if sk == "stretch":
                bud = bd.get("target") or 0.0
                if abs(bud - got) > 0.5: put(mix_full, bud - got, "balance", "balance of the budget number")
                if abs((T or 0.0) - bud) > 0.5: put(mix_full, (T or 0.0) - bud, "stretch", "the stretch above the budget (new business and the payments industry)")
            else:
                resid = (T or 0.0) - got
                if abs(resid) > 0.5:
                    put(mix_full, resid, "balance", "balance of the %s number" % ("reviewed" if sk == "today" else sk))
            by_cat = {k: round(b26[k] + alloc.get(k, 0.0), 2) for k in CATS}
            # subcategories inside each category, pro rata to 2026; a category the cell did not have goes to the lead subcategory
            subs = []
            for k in CATS:
                if abs(by_cat[k]) < 0.5 and not base[(cc, line)][k]: continue
                s26 = base[(cc, line)][k]
                if s26 and sum(s26.values()) > 0:
                    w = norm(s26)
                    for sub, ww in w.items():
                        v = by_cat[k] * ww
                        subs.append({"cat": k, "sub": sub, "y2026": round(s26[sub], 2), "v": round(v, 2), "rm": rm_for(cc, line, k, sub)})
                else:
                    sub = LEAD_SUB.get(line, "Other")
                    subs.append({"cat": k, "sub": sub, "y2026": 0.0, "v": round(by_cat[k], 2), "rm": rm_for(cc, line, k, sub)})
            rp = {k: round(sum(x["v"] * x["rm"] for x in subs if x["cat"] == k), 2) for k in CATS}
            out["scen"][sk] = {"total": round(T or 0.0, 2), "by_cat": by_cat, "rp": rp, "subs": subs,
                               "sources": {k: sorted(v, key=lambda x: -abs(x["eur"])) for k, v in src.items()},
                               "balance": round(resid, 2)}
        cells_out.append(out)

    # ---- roll-ups: group and countries, by category; retail/other flat ----
    def zero(): return {k: 0.0 for k in CATS}
    group = {"y2026": zero(), "rp_2026": zero()}
    for sk in list(scen) + ["today"]:
        group[sk] = zero(); group["rp_" + sk] = zero()
    countries = OrderedDict()
    for ent in entity_order:
        cc = bmap["countries"].get(ent)
        if cc: countries[cc] = {"code": cc, "country": name.get(cc, ent.replace("Printec ", "")), "entity": ent,
                                "y2026": zero(), "rp_2026": zero(), "flat": zero(), "flat_rp": zero(), "subs": [], "cells": []}
    for co in cells_out:
        cc = co["code"]
        if cc not in countries:
            countries[cc] = {"code": cc, "country": co["country"], "entity": entity_of.get(cc, "Printec " + co["country"]),
                             "y2026": zero(), "rp_2026": zero(), "flat": zero(), "flat_rp": zero(), "subs": [], "cells": []}
        C = countries[cc]
        for sk in list(scen) + ["today"]:
            C.setdefault(sk, zero()); C.setdefault("rp_" + sk, zero())
        for k in CATS:
            C["y2026"][k] += co["y2026"][k]; group["y2026"][k] += co["y2026"][k]
            r26 = sum(base_rp[(cc, co["line"])][k].values()); C["rp_2026"][k] += r26; group["rp_2026"][k] += r26
            for sk in list(scen) + ["today"]:
                C[sk][k] += co["scen"][sk]["by_cat"][k]; group[sk][k] += co["scen"][sk]["by_cat"][k]
                C["rp_" + sk][k] += co["scen"][sk]["rp"][k]; group["rp_" + sk][k] += co["scen"][sk]["rp"][k]
        C["cells"].append(co)
    flat_tot = zero(); flat_rp_tot = zero()
    for cc, cats in flat.items():
        if cc not in countries: continue
        for k, subs in cats.items():
            for sub, (rv, rp) in subs.items():
                countries[cc]["flat"][k] += rv; countries[cc]["flat_rp"][k] += rp; flat_tot[k] += rv; flat_rp_tot[k] += rp
                countries[cc]["subs"].append({"cat": k, "sub": sub, "kind": flat_kind.get(sub, "other"), "y2026": round(rv, 2),
                                              "floor": round(rv, 2), "budget": round(rv, 2), "upside": round(rv, 2), "today": round(rv, 2), "flat": True,
                                              "rp_budget": round(rp, 2), "rm": (rp / rv if rv else 0.0), "flat": True})
    # subcategory rows per country (banking cells), merged by cat+sub across lines
    for cc, C in countries.items():
        merged = OrderedDict()
        for co in C["cells"]:
            for sk in list(scen) + ["today"]:
                for x in co["scen"][sk]["subs"]:
                    key = (x["cat"], x["sub"])
                    m = merged.setdefault(key, {"cat": x["cat"], "sub": x["sub"], "kind": "banking", "y2026": 0.0, "budget": 0.0, "today": 0.0, "rp_budget": 0.0, "lines": set(), **({} if single else {"floor": 0.0, "upside": 0.0}), **({"stretch": 0.0} if ST else {})})
                    m[sk] += x["v"]
                    if sk == "budget":
                        m["y2026"] += x["y2026"]; m["rp_budget"] += x["v"] * x["rm"]
                    m["lines"].add(co["line_label"])
        for m in merged.values():
            m["lines"] = sorted(m["lines"]); m["rm"] = (m["rp_budget"] / m["budget"]) if m["budget"] else 0.0
            for k in ("y2026", "budget", "today", "rp_budget") + (() if single else ("floor", "upside")) + (("stretch",) if ST else ()): m[k] = round(m[k], 2)
            C["subs"].append(m)
        C["subs"].sort(key=lambda x: (CATS.index(x["cat"]), -x["budget"]))
        # totals incl. flat
        for sk in ["y2026"] + list(scen) + ["today"]:
            C["total_" + sk] = round(sum(C[sk].values()) + sum(C["flat"].values()), 2)
        C["rp_total_budget"] = round(sum(C["rp_budget"].values()) + sum(C["flat_rp"].values()), 2)
        C["rp_total_2026"] = round(sum(C["rp_2026"].values()) + sum(C["flat_rp"].values()), 2)
        for k in ("y2026", "rp_2026", "flat", "flat_rp") + tuple(list(scen) + ["today"]) + tuple("rp_" + s2 for s2 in list(scen) + ["today"]):
            C[k] = {kk: round(vv, 2) for kk, vv in C[k].items()}
        # the written explanation, per category
        C["text"] = explain(C, cc)
        C["waterfall"] = {sk: [{"cat": k, "from": C["y2026"][k], "to": C[sk][k], "delta": round(C[sk][k] - C["y2026"][k], 2)} for k in CATS] for sk in list(scen) + ["today"]}
        for co in C["cells"]:
            for sk in co["scen"].values():
                sk.pop("subs", None)     # the country-level subs carry the layout; keep the cell light
    for sk in ["y2026"] + list(scen) + ["today"]:
        group["total_" + sk] = round(sum(group[sk].values()) + sum(flat_tot.values()), 2)
    group["flat"] = {k: round(v, 2) for k, v in flat_tot.items()}; group["flat_rp"] = {k: round(v, 2) for k, v in flat_rp_tot.items()}
    for k in list(group):
        if isinstance(group[k], dict): group[k] = {kk: round(vv, 2) for kk, vv in group[k].items()}
    group["waterfall"] = {sk: [{"cat": k, "from": group["y2026"][k], "to": group[sk][k], "delta": round(group[sk][k] - group["y2026"][k], 2)} for k in CATS] for sk in list(scen) + ["today"]}
    group["rp_total_2026"] = round(sum(group["rp_2026"].values()) + sum(flat_rp_tot.values()), 2)
    for sk in list(scen) + ["today"]:
        group["rp_total_" + sk] = round(sum(group["rp_" + sk].values()) + sum(flat_rp_tot.values()), 2)
    # by source (group), per scenario
    by_source = {}
    for sk in list(scen) + ["today"]:
        agg = defaultdict(lambda: zero())
        for co in cells_out:
            for k, lst in co["scen"][sk]["sources"].items():
                for x in lst: agg[x["kind"]][k] += x["eur"]
        # presentation follows the drivers (22/09/2026): work by existing staff sits inside the model's deals,
        # and the model's recurring growth inside the price and efficiency steps, split in their proportion
        if "capacity" in agg:
            for k in CATS: agg["model_deals"][k] += agg["capacity"][k]
            del agg["capacity"]
        if "model_recurring" in agg:
            tp = sum(agg["pricing"].values()) if "pricing" in agg else 0.0; te = sum(agg["efficiency"].values()) if "efficiency" in agg else 0.0
            wp = tp / (tp + te) if (tp + te) else 1.0
            for k in CATS:
                agg["pricing"][k] += agg["model_recurring"][k] * wp; agg["efficiency"][k] += agg["model_recurring"][k] * (1 - wp)
            del agg["model_recurring"]
        by_source[sk] = {kind: {k: round(v, 2) for k, v in d.items()} for kind, d in agg.items()}
    # reconciliation against the plan
    Tb = B["totals"]
    checks = {"budget": [group["total_budget"], Tb["commit"]], "y2026": [group["total_y2026"], Tb["rev_2026"]]}
    if ST: checks["stretch"] = [group["total_stretch"], ST["total"]]
    if not single: checks.update({"floor": [group["total_floor"], Tb["floor"]], "upside": [group["total_upside"], Tb["upside"]]})
    for k, (a, b) in checks.items():
        if abs(a - b) > 5:
            print("WARNING: %s categories %.2f vs plan %.2f" % (k, a, b), file=sys.stderr)
    rules = [r for r in RULES if not (single and ("closing facts" in r or "lowest acceptable" in r))]
    if single: rules = [r.replace("Hires go to services (to outsourcing when the line is managed services). ", "Revenue that existing staff deliver in place of the former hire actions goes to services (to outsourcing when the line is managed services). ") for r in rules]
    return {"generated": datetime.datetime.now().isoformat(timespec="seconds"), "built_from": B.get("generated"), "single_number": single,
            "categories": CATS, "cat_label": CAT_LABEL, "scenarios": list(scen) + ["today"],
            "rules": rules, "group": group, "by_source": by_source, "countries": list(countries.values()),
            "cells": cells_out, "checks": {k: {"categories": round(a, 2), "plan": round(b, 2)} for k, (a, b) in checks.items()},
            "profit_note": "Profit holds each country's 2026 margin by category and subcategory. The plan's own profit figures hold the margin by product line, so the two differ where the growth is heavier in hardware or in services than the line average."}

RULES = [
    "The 2027 numbers are the plan's numbers, unchanged. This view only splits them into the Excel's categories.",
    "A dated order or a share win is split the way the country's 2026 revenue for that subcategory already splits across hardware, software, services and outsourcing, so installation and first-year support land in services.",
    "The model's recurring growth goes to services and outsourcing in their 2026 ratio; its named-deal and market-outlook growth follows the full 2026 mix.",
    "Hires go to services (to outsourcing when the line is managed services). Capability decisions go to software. Price and efficiency steps go to services and outsourcing.",
    "The three closing facts in the lowest acceptable number are orders and follow the full mix.",
    "A combination with no 2026 revenue uses the group's 2026 mix for its product line; a first outsourcing contract is 100% outsourcing under ATM.",
    "Retail lines and budget rows with no product name (Other) are carried flat at their 2026 value inside their own category: hardware, software, services or outsourcing. There is no separate retail item.",
    "Profit holds the 2026 margin of each country, category and subcategory.",
]

def explain(C, cc):
    """Plain-English paragraph per category for one country: the 2026 to 2027 move and what produces it."""
    txt = {}
    for k in CATS:
        a, b = C["y2026"][k], C["budget"][k]
        if a < 0.5 and b < 0.5: continue
        srcs = defaultdict(float)
        for co in C["cells"]:
            for x in co["scen"]["budget"]["sources"].get(k, []):
                srcs[(x["kind"], x["label"], co["line_label"])] += x["eur"]
        items = sorted(srcs.items(), key=lambda x: -abs(x[1]))
        g = ("%+.1f%%" % ((b / a - 1) * 100)) if a else "new"
        head = "%s goes from %s to %s (%s)." % (CAT_LABEL[k], M(a), M(b), g) if a else "%s starts from nothing and reaches %s." % (CAT_LABEL[k], M(b))
        if a and a < 100000 and b > 3 * a: head = "%s goes from %s to %s, from a small base." % (CAT_LABEL[k], M(a), M(b))
        lead = []
        for (kind, label, line), eur in items[:5]:
            if abs(eur) < 20000: continue
            e = "EUR %+.2fm" % (eur / 1e6)
            if kind.startswith("model"): lead.append("%s in %s (%s)" % (label, line, e))
            elif kind == "balance": lead.append("the balance of the plan number in %s (%s)" % (line, e))
            elif kind == "closing_fact": lead.append("a closing fact in %s (%s)" % (line, e))
            elif kind in ("date", "share"): lead.append("its share of %s (%s)" % (label, e))
            elif kind == "capacity": lead.append("%s (%s)" % (label, e))
            else: lead.append("%s (%s)" % (label, e))
        rest = len([1 for (_, _, _), e in items[5:] if abs(e) >= 20000])
        body = (" It carries " + "; ".join(lead) + (("; and %d smaller items" % rest) if rest else "") + ".") if lead else ""
        txt[k] = head + body
    return txt

def write_xlsx(out, path, bmap):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "BGT 2027 by category"
    single = bool(out.get("single_number"))
    hdr = ["Country / Category / Subcategory", "Revenue 2026 (BGT)", "Lowest acceptable 2027", "Budget 2027", "Maximum 2027", "RP 2027 (budget)", "RM% 2027 (budget)", "Product lines (dashboard)"]
    g = lambda o, k: o.get(k, o.get("budget")) if isinstance(o, dict) else o
    ws.append(hdr)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="EFE7D6")
    bold = Font(bold=True); catf = Font(bold=True, color="5F594C")
    for C in out["countries"]:
        ws.append([C["entity"], C["total_y2026"], C.get("total_floor", C["total_budget"]), C["total_budget"], C.get("total_upside", C["total_budget"]), C["rp_total_budget"], (C["rp_total_budget"] / C["total_budget"]) if C["total_budget"] else 0, ""])
        for c in ws[ws.max_row]: c.font = bold
        for k in CATS:
            rows = [s for s in C["subs"] if s["cat"] == k]
            if not rows: continue
            rp = sum(s["rp_budget"] for s in rows); tb = sum(s["budget"] for s in rows)
            ws.append([k, sum(s["y2026"] for s in rows), sum(s.get("floor", s["budget"]) for s in rows), tb, sum(s.get("upside", s["budget"]) for s in rows), rp, (rp / tb) if tb else 0, ""])
            for c in ws[ws.max_row]: c.font = catf
            for s in rows:
                ws.append([s["sub"], s["y2026"], s.get("floor", s["budget"]), s["budget"], s.get("upside", s["budget"]), s["rp_budget"], s["rm"], ", ".join(s.get("lines", [])) if s["kind"] == "banking" else ("retail, carried flat" if s["kind"] == "retail" else "other, carried flat")])
    G = out["group"]
    ws.append(["TOTAL", G["total_y2026"], G.get("total_floor", G["total_budget"]), G["total_budget"], G.get("total_upside", G["total_budget"]), G["rp_total_budget"], (G["rp_total_budget"] / G["total_budget"]) if G["total_budget"] else 0, ""])
    for c in ws[ws.max_row]: c.font = bold
    ws.append([]); ws.append(["Rules"]); ws[ws.max_row][0].font = bold
    for r in out["rules"]: ws.append([r])
    ws.append([out["profit_note"]])
    for row in ws.iter_rows(min_row=2):
        for c in row[1:6]: c.number_format = "#,##0"
        row[6].number_format = "0.0%"
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF": ws.column_dimensions[col].width = 20
    ws.column_dimensions["G"].width = 14; ws.column_dimensions["H"].width = 48
    ws.freeze_panes = "B2"
    if single:
        ws.delete_cols(5); ws.delete_cols(3)     # no lowest acceptable, no maximum: one budget number
        if "stretch" in out.get("scenarios", []):
            G = out["group"]; ws.cell(row=1, column=ws.max_column + 1, value="Stretch 2027 (EUR 200.0m)").font = Font(bold=True)
            col = ws.max_column; r = 2
            for C in out["countries"]:
                ws.cell(row=r, column=col, value=C.get("total_stretch")); r += 1
                for k in CATS:
                    rows = [s for s in C["subs"] if s["cat"] == k]
                    if not rows: continue
                    ws.cell(row=r, column=col, value=sum(s.get("stretch", s["budget"]) for s in rows)); r += 1
                    for s in rows:
                        ws.cell(row=r, column=col, value=s.get("stretch", s["budget"])); r += 1
            ws.cell(row=r, column=col, value=G.get("total_stretch"))
            for rr in range(2, r + 1): ws.cell(row=rr, column=col).number_format = "#,##0"
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 22
    wb.save(path)

def main():
    root = sys.argv[1] if len(sys.argv) > 1 else ROOT
    out = compute(root)
    bdir = os.path.join(root, "budget")
    save(os.path.join(bdir, "budget_2027_categories.json"), out)
    bmap = load(os.path.join(root, "weekly-intelligence", "references", "budget_map.json"))
    write_xlsx(out, os.path.join(bdir, "budget_2027_by_category.xlsx"), bmap)
    G = out["group"]
    fl = lambda k: G.get("floor", G["budget"])[k]; up = lambda k: G.get("upside", G["budget"])[k]
    print("group by category (EUR m)  2026 -> floor / budget / maximum" if not out.get("single_number") else "group by category (EUR m)  2026 -> budget")
    for k in CATS:
        print("  %-4s %7.2f -> %7.2f / %7.2f / %7.2f" % (k, G["y2026"][k] / 1e6, fl(k) / 1e6, G["budget"][k] / 1e6, up(k) / 1e6))
    print("  flat %7.2f" % (sum(G["flat"].values()) / 1e6))
    print("  total %6.2f -> %7.2f / %7.2f / %7.2f   profit (budget) %.2f vs 2026 %.2f" % (G["total_y2026"] / 1e6, G.get("total_floor", G["total_budget"]) / 1e6, G["total_budget"] / 1e6, G.get("total_upside", G["total_budget"]) / 1e6, G["rp_total_budget"] / 1e6, G["rp_total_2026"] / 1e6))
    print("checks:", out["checks"])
    print("wrote budget/budget_2027_categories.json and budget/budget_2027_by_category.xlsx")

if __name__ == "__main__":
    main()
